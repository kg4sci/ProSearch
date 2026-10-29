import json
from pathlib import Path
from collections import OrderedDict, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment

RECORD_TYPE_FILLS = {
    "field": PatternFill("solid", fgColor="DDEBF7"),      # 浅蓝
    "setup": PatternFill("solid", fgColor="FFF2CC"),      # 浅黄
    "paragraph": PatternFill("solid", fgColor="E2F0D9"),  # 浅绿
}

FIELD_ROLE_FILLS = {
    "fixed": PatternFill("solid", fgColor="EDEDED"),      # 很浅灰
    "explore": PatternFill("solid", fgColor="FCE4D6"),    # 很浅橙粉
}

INPUT_FILE = Path("data/audit_sample_data.jsonl")
OUTPUT_DIR = Path("data/manual_audit")


def load_jsonl(path: Path):
    records = []
    with path.open("r", encoding="utf-8") as f:
        for line_no, line in enumerate(f, 1):
            line = line.strip()
            if not line:
                continue
            try:
                records.append(json.loads(line))
            except json.JSONDecodeError as e:
                raise ValueError(f"JSONL 解析失败: {path}, line {line_no}") from e
    return records


def sanitize_name(name: str) -> str:
    if not name:
        return "unknown"
    return (
        str(name)
        .strip()
        .replace("/", "_")
        .replace("\\", "_")
        .replace(" ", "_")
    )


def dump_excel_sheet(wb, sheet_name, rows, fieldnames):
    ws = wb.create_sheet(title=sheet_name)

    # 表头
    for col_idx, field in enumerate(fieldnames, start=1):
        cell = ws.cell(row=1, column=col_idx, value=field)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 数据
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, field in enumerate(fieldnames, start=1):
            value = row.get(field, "")
            ws.cell(row=row_idx, column=col_idx, value=value)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(fieldnames))}{max(len(rows) + 1, 2)}"

    width_overrides = {
        "sample_id": 28,
        "candidate_id": 28,
        "paragraph_id": 14,
        "domain": 18,
        "type": 22,
        "type_id": 10,
        "task_family": 14,
        "query": 42,
        "constraints_text": 60,
        "explore_field": 22,
        "paragraph_head": 30,
        "paragraph_doi": 32,
        "paragraph_text": 110,
        "record_type": 12,
        "field_name": 32,
        "field_role": 12,
        "constraint_surface": 42,
        "model_value": 24,
        "model_judge": 16,
        "model_reason": 60,
        "human_label": 18,
        "notes": 28,
        "faithfulness": 12,
        "clarity": 10,
        "retrieval_usefulness": 18,
        "decision": 12,
        "paragraph_label": 16,
        "sample_label": 16,
        "boundary_focus_type": 22,
        "judge_paragraph_verdict_existing": 24,
        "model": 18,
        "setup_idx": 10,
        "setup_id": 14,
    }

    for col_idx, field in enumerate(fieldnames, start=1):
        if field in width_overrides:
            width = width_overrides[field]
        else:
            max_len = len(str(field))
            for row in rows[:200]:
                val = row.get(field, "")
                val_str = "" if val is None else str(val)
                max_len = max(max_len, min(len(val_str), 50))
            width = min(max_len + 2, 50)

        ws.column_dimensions[get_column_letter(col_idx)].width = width

    wrap_fields = {
        "query",
        "constraints_text",
        "paragraph_text",
        "constraint_surface",
        "model_reason",
        "notes",
    }

    for col_idx, field in enumerate(fieldnames, start=1):
        if field in wrap_fields:
            for row_idx in range(2, len(rows) + 2):
                ws.cell(row=row_idx, column=col_idx).alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

    for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1):
        for cell in row:
            if cell.alignment == Alignment():
                cell.alignment = Alignment(vertical="top")

        # 按 record_type / field_role 上色
    field_to_col = {field: idx for idx, field in enumerate(fieldnames, start=1)}

    record_type_col = field_to_col.get("record_type")
    field_role_col = field_to_col.get("field_role")
    field_name_col = field_to_col.get("field_name")

    for row_idx, row in enumerate(rows, start=2):
        record_type = row.get("record_type", "")
        field_role = row.get("field_role", "")

        # 1) record_type：整行浅色
        if record_type in RECORD_TYPE_FILLS:
            fill = RECORD_TYPE_FILLS[record_type]
            for col_idx in range(1, len(fieldnames) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill

        # 2) field_role：只给关键列更深一点的提示色，避免太花
        if record_type == "field" and field_role in FIELD_ROLE_FILLS:
            role_fill = FIELD_ROLE_FILLS[field_role]

            if field_role_col:
                ws.cell(row=row_idx, column=field_role_col).fill = role_fill
            if field_name_col:
                ws.cell(row=row_idx, column=field_name_col).fill = role_fill

def format_constraint(c):
    field = c.get("field", "")
    op = c.get("op", "")
    if op == "=":
        return f"{field} = {c.get('value', '')}"
    elif op == "!=":
        return f"{field} != {c.get('value', '')}"
    elif op == "between":
        r = c.get("range", {}) or {}
        lower = r.get("lower", "")
        upper = r.get("upper", "")
        unit = r.get("unit", "")
        return f"{field} between {lower}-{upper} {unit}".strip()
    elif op in {">", ">=", "<", "<="}:
        return f"{field} {op} {c.get('value', '')}"
    else:
        if "surface" in c:
            return f"{field} {op} {c['surface']}"
        return json.dumps(c, ensure_ascii=False)


def build_constraints_text(constraints, explore_field=None):
    parts = []
    for c in constraints or []:
        parts.append(format_constraint(c))
    return " ;\n ".join(parts)


def constraint_surface_map(constraints):
    m = {}
    for c in constraints or []:
        field = c.get("field")
        if field:
            m[field] = format_constraint(c)
    return m


def safe_join_values(v):
    if v is None:
        return ""
    if isinstance(v, list):
        vals = []
        for x in v:
            if isinstance(x, dict):
                vals.append(str(x.get("value", "")))
            else:
                vals.append(str(x))
        return " | ".join(vals)
    return str(v)


def extract_judge_info(rec):
    s = rec.get("llm_judge_summary") or {}
    full = rec.get("llm_judge_record") or {}

    constraints = s.get("constraints")
    if constraints is None:
        constraints = full.get("constraints", [])

    explore_field = s.get("explore_field")
    if explore_field is None:
        explore_field = full.get("explore_field")

    setup_judgments = s.get("setup_judgments")
    if setup_judgments is None:
        parsed_output = full.get("parsed_output", {}) or {}
        setup_judgments = parsed_output.get("setup_judgments", [])

    paragraph_verdict = s.get("paragraph_verdict")
    paragraph_reason = s.get("paragraph_reason")
    if paragraph_verdict is None or paragraph_reason is None:
        parsed_output = full.get("parsed_output", {}) or {}
        paragraph_verdict = parsed_output.get("paragraph_verdict")
        paragraph_reason = parsed_output.get("paragraph_reason")

    return {
        "model": s.get("model") or full.get("model", ""),
        "constraints": constraints or [],
        "explore_field": explore_field,
        "setup_judgments": setup_judgments or [],
        "paragraph_verdict": paragraph_verdict,
        "paragraph_reason": paragraph_reason,
    }


def build_query_audit_rows(records):
    seen = OrderedDict()

    for rec in records:
        candidate_id = rec["candidate_id"]
        judge_info = extract_judge_info(rec)
        constraints = judge_info["constraints"]
        explore_field = judge_info["explore_field"]

        row = {
            "candidate_id": candidate_id,
            "domain": rec.get("domain", ""),
            "type": rec.get("type", ""),
            "type_id": rec.get("type_id", ""),
            "task_family": rec.get("task_family", ""),
            "query": rec.get("query", ""),
            "constraints_text": build_constraints_text(constraints, explore_field),
            "explore_field": explore_field or "",
            "faithfulness": "",
            "clarity": "",
            "retrieval_usefulness": "",
            "decision": "",
            "notes": "",
        }
        seen[candidate_id] = row

    return list(seen.values())


def build_relevance_long_rows(records):
    rows = []

    for rec in records:
        candidate_id = rec["candidate_id"]
        paragraph_id = str(rec["paragraph_id"])
        sample_id = f"{candidate_id}__{paragraph_id}"

        judge_info = extract_judge_info(rec)
        constraints = judge_info["constraints"]
        explore_field = judge_info["explore_field"]
        setup_judgments = judge_info["setup_judgments"]
        paragraph_verdict = judge_info["paragraph_verdict"]
        paragraph_reason = judge_info["paragraph_reason"]
        model = judge_info["model"]

        c_surface_map = constraint_surface_map(constraints)

        base_common = {
            "sample_id": sample_id,
            "candidate_id": candidate_id,
            "paragraph_id": paragraph_id,
            "domain": rec.get("domain", ""),
            "type": rec.get("type", ""),
            "type_id": rec.get("type_id", ""),
            "task_family": rec.get("task_family", ""),
            "paragraph_label": rec.get("paragraph_label", ""),
            "sample_label": rec.get("sample_label", ""),
            "boundary_focus_type": rec.get("boundary_focus_type", ""),
            "judge_paragraph_verdict_existing": rec.get("judge_paragraph_verdict", ""),
            "query": rec.get("query", ""),
            "constraints_text": build_constraints_text(constraints, explore_field),
            "explore_field": explore_field or "",
            "model": model or "",
        }

        for s_idx, setup in enumerate(setup_judgments, start=1):
            setup_id = setup.get("setup_id", f"setup_{s_idx}")
            fixed_fields = setup.get("fixed_field_judgments", []) or []
            explore_j = setup.get("explore_field_judgment") or {}

            # fixed field rows
            for fj in fixed_fields:
                field_name = fj.get("field", "")
                row = dict(base_common)
                row.update({
                    "setup_idx": s_idx,
                    "setup_id": setup_id,
                    "record_type": "field",
                    "field_name": field_name,
                    "field_role": "fixed",
                    "constraint_surface": c_surface_map.get(field_name, ""),
                    "model_value": safe_join_values(fj.get("value")),
                    "model_judge": fj.get("status", ""),
                    "model_reason": fj.get("evidence", ""),
                    "human_label": "",
                    "notes": "",
                })
                rows.append(row)

            # explore field row
            if explore_field:
                row = dict(base_common)
                row.update({
                    "setup_idx": s_idx,
                    "setup_id": setup_id,
                    "record_type": "field",
                    "field_name": explore_field,
                    "field_role": "explore",
                    "constraint_surface": f"EXPLORE: {explore_field}",
                    "model_value": safe_join_values(explore_j.get("values")),
                    "model_judge": explore_j.get("status", ""),
                    "model_reason": "",
                    "human_label": "",
                    "notes": "",
                })
                rows.append(row)

            # setup-level row
            row = dict(base_common)
            row.update({
                "setup_idx": s_idx,
                "setup_id": setup_id,
                "record_type": "setup",
                "field_name": "",
                "field_role": "",
                "constraint_surface": "",
                "model_value": "",
                "model_judge": setup.get("setup_verdict", ""),
                "model_reason": setup.get("setup_reason", ""),
                "human_label": "",
                "notes": "",
            })
            rows.append(row)

        # paragraph-level row
        row = dict(base_common)
        row.update({
            "setup_idx": 0,
            "setup_id": "",
            "record_type": "paragraph",
            "field_name": "",
            "field_role": "",
            "constraint_surface": "",
            "model_value": "",
            "model_judge": paragraph_verdict or "",
            "model_reason": paragraph_reason or "",
            "human_label": "",
            "notes": "",
        })
        rows.append(row)

    def sort_key(r):
        record_order = {"field": 0, "setup": 1, "paragraph": 2}
        field_role_order = {"fixed": 0, "explore": 1, "": 2}
        return (
            r["sample_id"],
            int(r["setup_idx"]),
            record_order.get(r["record_type"], 99),
            field_role_order.get(r["field_role"], 99),
            r["field_name"],
        )

    rows.sort(key=sort_key)
    return rows


def build_paragraph_rows(records):
    """
    段落单独存。
    这里按 sample_id 存一行，方便和 manual_audit 里的 sample_id 直接对应。
    """
    rows = []
    seen = set()

    for rec in records:
        candidate_id = rec["candidate_id"]
        paragraph_id = str(rec["paragraph_id"])
        sample_id = f"{candidate_id}__{paragraph_id}"
        if sample_id in seen:
            continue
        seen.add(sample_id)

        para_meta = rec.get("paragraph_meta", {}) or {}

        rows.append({
            "sample_id": sample_id,
            "candidate_id": candidate_id,
            "paragraph_id": paragraph_id,
            "domain": rec.get("domain", ""),
            "type": rec.get("type", ""),
            "type_id": rec.get("type_id", ""),
            "task_family": rec.get("task_family", ""),
            "paragraph_head": para_meta.get("head", ""),
            "paragraph_doi": para_meta.get("doi", ""),
            "paragraph_text": rec.get("paragraph_text", ""),
        })

    rows.sort(key=lambda r: (r["candidate_id"], r["paragraph_id"]))
    return rows


def save_workbook(path: Path, sheets):
    """
    sheets: [(sheet_name, rows, fieldnames), ...]
    """
    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)

    for sheet_name, rows, fieldnames in sheets:
        dump_excel_sheet(wb, sheet_name, rows, fieldnames)

    wb.save(path)


def group_records_by_domain(records):
    grouped = defaultdict(list)
    for rec in records:
        domain = sanitize_name(rec.get("domain", "unknown"))
        grouped[domain].append(rec)
    return grouped


def main():
    records = load_jsonl(INPUT_FILE)
    print(f"读取 enriched audit 样本: {len(records)}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    grouped = group_records_by_domain(records)

    query_fields = [
        "candidate_id",
        "domain",
        "type",
        "type_id",
        "task_family",
        "constraints_text",
        "explore_field",
        "query",
        "faithfulness",
        "clarity",
        "retrieval_usefulness",
        "decision",
        "notes",
    ]

    # 注意：这里已经去掉 paragraph_text
    relevance_fields = [
        "sample_id",
        "candidate_id",
        "paragraph_id",
        "domain",
        "type",
        "type_id",
        "task_family",
        "paragraph_label",
        "sample_label",
        "boundary_focus_type",
        "judge_paragraph_verdict_existing",
        "query",
        "constraints_text",
        "explore_field",
        "model",
        "setup_idx",
        "setup_id",
        "record_type",
        "field_name",
        "field_role",
        "constraint_surface",
        "model_value",
        "model_judge",
        "model_reason",
        "human_label",
        "notes",
    ]

    paragraph_fields = [
        "sample_id",
        "candidate_id",
        "paragraph_id",
        "domain",
        "type",
        "type_id",
        "task_family",
        "paragraph_head",
        "paragraph_doi",
        "paragraph_text",
    ]

    total_query_rows = 0
    total_relevance_rows = 0
    total_paragraph_rows = 0

    for domain, domain_records in sorted(grouped.items()):
        query_rows = build_query_audit_rows(domain_records)
        relevance_rows = build_relevance_long_rows(domain_records)
        paragraph_rows = build_paragraph_rows(domain_records)

        manual_audit_path = OUTPUT_DIR / f"manual_audit__{domain}.xlsx"
        paragraph_path = OUTPUT_DIR / f"paragraph_text__{domain}.xlsx"

        save_workbook(
            manual_audit_path,
            sheets=[
                ("query_audit", query_rows, query_fields),
                ("relevance_long", relevance_rows, relevance_fields),
            ]
        )

        save_workbook(
            paragraph_path,
            sheets=[
                ("paragraphs", paragraph_rows, paragraph_fields),
            ]
        )

        print(f"[{domain}] 已写出标注文件: {manual_audit_path}")
        print(f"[{domain}] 已写出段落文件: {paragraph_path}")
        print(f"[{domain}] query_audit 行数: {len(query_rows)}")
        print(f"[{domain}] relevance_long 行数: {len(relevance_rows)}")
        print(f"[{domain}] paragraphs 行数: {len(paragraph_rows)}")

        total_query_rows += len(query_rows)
        total_relevance_rows += len(relevance_rows)
        total_paragraph_rows += len(paragraph_rows)

    print("========== 汇总 ==========")
    print(f"领域数: {len(grouped)}")
    print(f"query_audit 总行数: {total_query_rows}")
    print(f"relevance_long 总行数: {total_relevance_rows}")
    print(f"paragraphs 总行数: {total_paragraph_rows}")


if __name__ == "__main__":
    main()
# python scripts/stage1_relevance/build_sample_data_excel.py
