import json
from typing import Dict, List, Any, Optional
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SYSTEM_PROMPT = """
You are a specialist in scientific setup-level constraint judgment.
"""

REPRODUCTION_PROMPT = """
You are judging whether a scientific paragraph contains at least one experimental setup that satisfies a structured specification.

Treat the paragraph as describing one or more possible setups. Your task is to decide whether there exists at least one concrete setup in the paragraph that matches all fixed constraints. All satisfied fields must belong to the same setup. Do not combine evidence from different setups.

Work through the paragraph in four steps.

Step 1. Identify the setup units.
Read the paragraph and separate it into setup-level experimental configurations. A new setup should be created only when the paragraph clearly shifts to a different system, composition, condition set, treatment path, reaction setting, measurement configuration, material-role assignment, or other distinct experimental configuration. Do not split one setup just because the paragraph repeats it, comments on it, or reports outcomes from it. Do not create a new setup from results alone unless the paragraph clearly introduces a different experimental configuration. If the paragraph does not contain any plausible setup that can be judged separately, return no setup judgments and decide the paragraph verdict at the end.

Step 2. Judge each fixed field within the same setup.
For each plausible setup, judge every fixed field inside that setup only. Do not judge fields at the whole-paragraph level, and do not mix evidence across setups.

A field does not need to appear in exactly the same words as the structured constraint. You may use short, evidence-grounded scientific reasoning when the mapping is direct and well supported. This includes standard terminology mapping, conventional scientific interpretation, direct recovery of a field value from a clearly described process, and exact unit conversion. However, do not guess, do not fill in missing information by speculation, and do not use long chains of inference. If the value is only weakly hinted at, mark it as missing.

For each fixed field, assign one of three statuses: match, violate, or missing.

For "=" constraints, use match when the constrained value is directly supported or can be directly recovered from the same setup by short evidence-grounded reasoning. Use violate only when the paragraph clearly supports an incompatible value for that same setup. Otherwise use missing.

For "!=" constraints, use violate when the forbidden value is directly supported or can be directly recovered from the same setup by short evidence-grounded reasoning. Use match when the field is present and the setup clearly supports a different value. If the field is absent or too weakly supported to judge, use missing.

Step 3. Judge numeric fields after exact unit conversion.
For numeric constraints, first apply exact unit conversion whenever the units are compatible. Then compare the normalized values within the same setup.

Use a simple overlap rule:
- If the constrained value or range overlaps with the paragraph value or range, mark match.
- If they are clearly disjoint, mark violate.
- If the paragraph gives only vague, approximate, partial, or non-comparable numeric information, mark missing.

Examples:
- Constraint = 5 mM; paragraph = 5 mM → match
- Constraint = 5 mM; paragraph = 3–7 mM → match
- Constraint = 3–7 mM; paragraph = 5 mM → match
- Constraint = 3–7 mM; paragraph = 2–8 mM → match
- Constraint = 3–7 mM; paragraph = 8–10 mM → violate
- Constraint = 0.5 h; paragraph = 1800 s → match
- Constraint = 3–7 mM; paragraph = about several mM → missing
- Constraint = 700 °C; paragraph = calcined at 400–800 °C for 1 h → match

When a field is match or violate, return one decisive supported value and one short evidence span from the paragraph. When a field is missing, return null for both value and evidence.

Step 4. Derive the setup verdict and then the paragraph verdict.
A setup is satisfy only when all fixed fields are match within that same setup.
A setup is uncertain when no fixed field is violate, at least one fixed field is match, and at least one remaining fixed field is missing.
A setup is violate when at least one fixed field is violate, or when none of the constrained fields is matched for that setup.

After judging all setups:
- The paragraph is satisfy if at least one setup is satisfy.
- The paragraph is uncertain if no setup is satisfy but at least one setup is uncertain.
- The paragraph is violate if all judged setups are violate, or if the paragraph does not contain any plausible setup that can be judged separately.

Return JSON only in the following format:

{{
  "setup_judgments": [
    {{
      "setup_id": "setup_1",
      "fixed_field_judgments": [
        {{
          "field": "field_name",
          "status": "match",
          "value": "supported value",
          "evidence": "short evidence span"
        }},
        {{
          "field": "another_field",
          "status": "missing",
          "value": null,
          "evidence": null
        }}
      ],
      "setup_verdict": "uncertain",
      "setup_reason": "short reason"
    }}
  ],
  "paragraph_verdict": "uncertain",
  "paragraph_reason": "short reason"
}}

If the paragraph does not contain any plausible setup that can be judged separately, return "setup_judgments": [] and still provide a paragraph-level verdict and reason.

Keep setup_reason and paragraph_reason brief, diagnostic, and grounded in the same evidence standard used for the field judgments.

Structured constraints:
{constraints_json}

Relevant field definitions:
{field_definitions_json}

Original paragraph:
{paragraph_text}
"""

EXPLORATION_PROMPT = """
You are judging whether a scientific paragraph contains at least one experimental setup that satisfies a structured specification and also provides the value of an explore field.

Treat the paragraph as describing one or more possible setups. Your task is to decide whether there exists at least one concrete setup in the paragraph that matches all fixed constraints and provides the explore field for that same setup. All fixed fields and the explore field must belong to the same setup. Do not combine evidence from different setups.

Work through the paragraph in four steps.

Step 1. Identify the setup units.
Read the paragraph and separate it into setup-level experimental configurations. A new setup should be created only when the paragraph clearly shifts to a different system, composition, condition set, treatment path, reaction setting, measurement configuration, material-role assignment, or other distinct experimental configuration. Do not split one setup just because the paragraph repeats it, comments on it, or reports outcomes from it. Do not create a new setup from results alone unless the paragraph clearly introduces a different experimental configuration. If the paragraph does not contain any plausible setup that can be judged separately, return no setup judgments and decide the paragraph verdict at the end.

Step 2. Judge each fixed field within the same setup.
For each plausible setup, judge every fixed field inside that setup only. Do not judge fields at the whole-paragraph level, and do not mix evidence across setups.

A field does not need to appear in exactly the same words as the structured constraint. You may use short, evidence-grounded scientific reasoning when the mapping is direct and well supported. This includes standard terminology mapping, conventional scientific interpretation, direct recovery of a field value from a clearly described process, and exact unit conversion. However, do not guess, do not fill in missing information by speculation, and do not use long chains of inference. If the value is only weakly hinted at, mark it as missing.

For each fixed field, assign one of three statuses: match, violate, or missing.

For "=" constraints, use match when the constrained value is directly supported or can be directly recovered from the same setup by short evidence-grounded reasoning. Use violate only when the paragraph clearly supports an incompatible value for that same setup. Otherwise use missing.

For "!=" constraints, use violate when the forbidden value is directly supported or can be directly recovered from the same setup by short evidence-grounded reasoning. Use match when the field is present and the setup clearly supports a different value. If the field is absent or too weakly supported to judge, use missing.

Step 3. Judge numeric fields and then the explore field.
For numeric constraints, first apply exact unit conversion whenever the units are compatible. Then compare the normalized values within the same setup.

Use a simple overlap rule:
- If the constrained value or range overlaps with the paragraph value or range, mark match.
- If they are clearly disjoint, mark violate.
- If the paragraph gives only vague, approximate, partial, or non-comparable numeric information, mark missing.

Examples:
- Constraint = 5 mM; paragraph = 5 mM → match
- Constraint = 5 mM; paragraph = 3–7 mM → match
- Constraint = 3–7 mM; paragraph = 5 mM → match
- Constraint = 3–7 mM; paragraph = 2–8 mM → match
- Constraint = 3–7 mM; paragraph = 8–10 mM → violate
- Constraint = 0.5 h; paragraph = 1800 s → match
- Constraint = 3–7 mM; paragraph = about several mM → missing
- Constraint = 700 °C; paragraph = calcined at 400–800 °C for 1 h → match


When a fixed field is match or violate, return one decisive supported value and one short evidence span from the paragraph. When a fixed field is missing, return null for both value and evidence.

Now judge the explore field for that same setup. The explore field counts as present when the paragraph either explicitly states its value or provides enough evidence to recover the value by short, direct, evidence-grounded scientific reasoning without guesswork. The field does not need to be named in exactly the same words as the schema. What matters is whether a knowledgeable reader could directly recover a supported value from the setup as described.

Use "present" when the explore field value is explicitly given or directly recoverable from the setup.
Use "missing" when the paragraph does not provide enough support to recover a value confidently, when only weak hints are present, or when recovering the value would require guesswork or multi-step speculation.

If the explore field is present, return the supported value or values with short evidence spans. If it is missing, return an empty list.

Step 4. Derive the setup verdict and then the paragraph verdict.
A setup is satisfy only when all fixed fields are match and the explore field is present within that same setup.
A setup is uncertain when no fixed field is violate, at least one fixed field is match, but some fixed field is missing or the explore field is missing.
A setup is violate when at least one fixed field is violate, or when none of the constrained fields is matched for that setup.

After judging all setups:
- The paragraph is satisfy if at least one setup is satisfy.
- The paragraph is uncertain if no setup is satisfy but at least one setup is uncertain.
- The paragraph is violate if all judged setups are violate, or if the paragraph does not contain any plausible setup that can be judged separately.

Return JSON only in the following format:

{{
  "setup_judgments": [
    {{
      "setup_id": "setup_1",
      "fixed_field_judgments": [
        {{
          "field": "field_name",
          "status": "match",
          "value": "supported value",
          "evidence": "short evidence span"
        }}
      ],
      "explore_field_judgment": {{
        "field": "explore_field_name",
        "status": "present",
        "values": [
          {{
            "value": "supported value",
            "evidence": "short evidence span"
          }}
        ]
      }},
      "setup_verdict": "satisfy",
      "setup_reason": "short reason"
    }}
  ],
  "paragraph_verdict": "satisfy",
  "paragraph_reason": "short reason"
}}

If the paragraph does not contain any plausible setup that can be judged separately, return "setup_judgments": [] and still provide a paragraph-level verdict and reason.

Keep setup_reason and paragraph_reason brief, diagnostic, and grounded in the same evidence standard used for the field judgments.

Explore field:
{explore_field}

Structured constraints:
{constraints_json}

Relevant field definitions:
{field_definitions_json}

Original paragraph:
{paragraph_text}
"""


JUDGE_FIELD_DEFINITIONS = {
    "cell_culture": {
        "cultured_cells.cell_type": (
            "The cultured cell identity or identities in the setup. Use the core established identity needed for stable distinction. "
            "Obvious naming variants of the same identity may be treated as the same value, but biologically distinct identities should not be merged. "
            "For a constraint on one identity, mark match if that identity is explicitly present in the same setup, even if other identities are also present. "
            "Mark violate when a different clearly incompatible identity is explicitly supported and the constrained identity is not. Otherwise mark missing."
        ),

        "cell_source_type": (
            "The source category of the cultured cells, using coarse labels such as primary, cell line, immortalized, stem-cell-derived, or patient-derived. "
            "Use the most specific supported category. Explicit source descriptors such as primary, immortalized, patient-derived, or stem-cell-derived directly support the corresponding label. "
            "Established named cell lines directly support cell line unless the paragraph clearly indicates a different source category. "
            "Do not guess source type from weak or indirect context."
        ),

        "culture_medium.basal_medium": (
            "The basal culture medium used in the setup, such as DMEM, RPMI-1640, IMDM, MEM, Ham's F12, or DMEM/F12. "
            "Keep only the basal medium identity and ignore serum, antibiotics, additives, supplements, supplier information, concentration, and strength unless they are part of the medium name itself. "
            "Standard naming variants of the same basal medium may be treated as the same value. "
            "Do not infer a basal medium from complete commercial medium names unless the basal medium itself is explicitly stated."
        ),

        "culture_medium.serum.identity": (
            "The serum identity used in the culture medium, such as fetal bovine serum, fetal calf serum, bovine calf serum, horse serum, or human serum. "
            "Keep only the serum identity and ignore concentration, heat inactivation, supplier information, and similar details that do not change serum type. "
            "Standard naming variants or direct standard subtype relations may be treated as the same value when scientifically clear. "
            "Do not treat proteins, serum replacements, or serum substitutes as serum identity, and do not infer serum identity when it is not stated."
        ),

        "culture_medium.serum.concentration": (
            "The serum concentration used in the culture medium, such as 10%, 2%, 0-10%, or serum-free. "
            "Use only values clearly tied to serum in the same setup. Serum-free or explicit absence of serum may be treated as 0%. "
            "Do not use concentrations for non-serum components. "
            "If multiple serum concentrations are mentioned, use the one explicitly linked to the active setup being judged."
        ),

        "culture_substrate_or_coating": (
            "The material, matrix, coating, surface, or named engineered substrate used for cell attachment, support, coating, or embedding. "
            "Keep only the specific substrate or coating identity, and ignore concentrations, dimensions, pore sizes, vendor names, and procedural wording. "
            "Do not use generic support words such as plate, well, dish, insert, scaffold, membrane, or film alone unless a specific material is also stated. "
            "When both a generic support term and a specific material are present, return the specific material or coating."
        ),

        "seeding_density": (
            "The initial seeding, plating, or inoculation amount used to start the culture. "
            "Use the stated cell number together with its denominator or culture unit when available. "
            "Counts without a denominator may be used only when the paragraph clearly states that they were seeded, plated, or inoculated as the starting amount. "
            "If multiple counts or densities are mentioned, use the one explicitly tied to the initial setup. "
            "Do not use later culture density or other non-initial values."
        ),

        "culture_mode": (
            "The organizational culture format of the setup, using coarse labels such as 2D, 2.5D, 3D, spheroid, organoid, suspension, chip, Transwell, ALI, co-culture, or monoculture. "
            "Use explicit culture descriptors or very clear local evidence. Terms such as co-culture directly support co-culture, and terms such as monolayer may support 2D when they clearly describe culture format. "
            "For a constraint of the form culture_mode != X, mark violate only when the setup explicitly supports X; mark match when it explicitly supports a different culture mode; otherwise mark missing. "
            "Do not infer culture mode from vessel type alone."
        ),

        "differentiation_state": (
            "The differentiation status of the cultured cells, restricted to differentiated, undifferentiated, dedifferentiated, or partially differentiated. "
            "Match semantically equivalent wording or a clearly stated experimental phase when the meaning is direct and well supported. "
            "Do not infer differentiation state from treatment, induction, maturation, culture duration, or routine stem-cell maintenance alone unless the paragraph explicitly ties them to differentiation status."
        ),

        "confluence": (
            "The confluence value, range, or qualitative confluence state stated for the setup, such as 80%, 80-85%, near confluence, or full confluence. "
            "Decimal fractions may be interpreted as percentages when the paragraph clearly uses them as confluence. "
            "Ignore approximation markers and procedural wording when they do not change the core meaning. "
            "Do not use time durations, cells-per-area measurements, or generic density expressions unless they are explicitly presented as confluence."
        ),
    },

    "electrochemistry": {
        "electrolytes.chemical": (
            "The electrolyte-component identity or identities explicitly used in this electrochemical setup, such as salts, acids, bases, solvents, ionic liquids, added water, or buffers. "
            "Use the core chemical identity needed for stable distinction, and do not include concentration, pH, ratio, purity, or supplier information as part of identity. "
            "Standard naming variants, abbreviation versus full-name variants, and obvious formula-style variants may be treated as the same value when they preserve the same chemical identity, but chemically distinct components should not be merged. "
            "A setup may contain multiple explicit electrolyte components; for a constraint on one component, mark match if that component is explicitly present, even if other components are also present. "
            "Mark violate only when a different clearly incompatible component is explicitly supported and the constrained component is not; otherwise mark missing."
        ),

        "electrolytes.concentration": (
            "The concentration or composition value explicitly associated with an electrolyte component in this setup, such as 0.1 M, 50 mM, 1 wt%, 10 vol%, 100 ppm, or 1:1 (v/v). "
            "Use only values clearly tied to the relevant electrolyte component in the same setup. "
            "Equivalent unit surface forms may be treated as the same value when they preserve the same concentration or composition meaning. "
            "Do not use generic mixture descriptions or nearby numerical values unless they are clearly the concentration or composition of the relevant electrolyte component."
        ),

        "reference_electrode": (
            "The reference-electrode identity explicitly used in this setup, such as Ag/AgCl, SCE, RHE, or Li/Li+. "
            "Use the core reference-electrode identity needed for stable distinction. Standard naming variants, abbreviation versus full-name variants, and obvious notation variants may be treated as the same value when the reference-electrode identity is unchanged. "
            "Preserve distinctions that matter for matching, including different reference families and explicitly stated quasi-reference or pseudo-reference status. "
            "Do not infer a reference electrode from potential notation, common practice, or missing context when it is not explicitly supported."
        ),

        "working_electrode": (
            "The working-electrode identity explicitly used in this setup, such as Pt, Au, glassy carbon, Bi(111), carbon cloth, or another clearly stated working-electrode material or named electrode system. "
            "Use the core identity needed for stable distinction, including named modifiers, coatings, deposited layers, supports, or crystal faces only when they are clearly part of the working-electrode identity. "
            "Standard naming variants may be treated as the same value when the identity is unchanged, but materially different electrodes should not be merged. "
            "Do not treat generic apparatus descriptions, cell hardware, or weak contextual hints as sufficient evidence for a specific working electrode."
        ),

        "target_species": (
            "The analyte, reactant, intermediate, or key electroactive species that is the main target of detection, oxidation, reduction, quantification, or mechanistic study in this setup. "
            "Use the target-species identity needed for stable distinction, such as O2, CO2, H2O2, glucose, or nitrate. "
            "Standard name versus formula variants may be treated as the same value when they preserve the same species identity, but different chemical species, oxidation states, charge states, complexes, or explicitly stated molecular forms should not be merged. "
            "Do not treat ordinary supporting electrolyte components, byproducts, or background species as target species unless the paragraph clearly supports them as the electrochemical target."
        ),

        "temperature": (
            "The experimental temperature condition explicitly used in this setup, such as 25 °C, 80 °C, room temperature, or 298 K. "
            "Use only values that clearly denote temperature. Equivalent Celsius and Kelvin expressions may be treated as compatible when the correspondence is clear. "
            "Do not treat drying conditions, water content, thermal processing labels, or electrolyte-conditioning descriptions as temperature unless they clearly specify the experimental temperature of the setup."
        ),

        "measurement_method": (
            "The electrochemical measurement method or technique explicitly used in this setup, such as cyclic voltammetry, chronoamperometry, electrochemical impedance spectroscopy, differential pulse voltammetry, square wave voltammetry, linear sweep voltammetry, or chronopotentiometry. "
            "Use the method identity needed for stable distinction. Standard full-name and abbreviation variants may be treated as the same value when the technique identity is unchanged. "
            "Preserve distinctions that matter for matching, including different technique families and clearly stated method variants that correspond to different techniques. "
            "Do not treat generic labels such as electrochemical measurement, electrochemical test, or electrochemical analysis as sufficient evidence when a specific method is required."
        ),

        "scan_rate": (
            "The scan rate explicitly used in a voltammetric measurement in this setup, such as 50 mV/s, 0.1 V/s, or 10-100 mV/s. "
            "Use only values clearly stated as scan rate. Equivalent unit surface forms may be treated as the same value when they preserve the same meaning. "
            "Do not use step potentials, pulse amplitudes, frequencies, sweep numbers, or other nearby electrochemical parameters as scan rate unless they are explicitly identified as scan rate."
        ),

        "applied_potential": (
            "The applied potential, fixed potential, or potential window explicitly used in this setup, such as -0.8 V, 0.2-1.0 V, or 1.2 V vs Ag/AgCl. "
            "Use the potential condition itself, including the reference basis when it is explicitly stated. Equivalent surface forms may be treated as the same value when they preserve the same potential meaning. "
            "Potential values with different explicit reference bases should not be treated as equivalent unless the paragraph directly supports the same basis. "
            "Do not use onset potentials, peak positions, overpotentials, or other derived electrochemical results as applied potential unless the paragraph clearly states them as the imposed potential condition."
        ),
    },

    "immunotechnology": {
        "protein_targets.target": (
            "The core protein target intentionally probed in this Western blot. Use the target identity itself, not antibody product descriptions, vendor information, residue-site notation, or target-form wording such as phospho, total, or cleaved. "
            "Standard naming variants that preserve the same target identity may be treated as the same value, but distinct proteins, isoforms, subunits, or family members should not be merged unless the paragraph clearly supports the same target identity. "
            "If the paragraph clearly assigns a protein a loading-control or normalization role, do not treat it as the target protein for that blot."
        ),

        "protein_targets.target_form": (
            "The probed target form or state in this Western blot, such as phospho, total, cleaved, precursor, mature, or activated. "
            "Use only forms clearly grounded in the paragraph. Obvious local variants such as p-, phospho-, or phosphorylated may support phospho, and total plus the target name may support total. "
            "Residue-site notation may support phospho when it clearly refers to a phosphorylated form of the target. "
            "Do not infer target form from nearby context alone when the paragraph does not clearly support it."
        ),

        "normalization_control": (
            "The protein explicitly used as a loading control or normalization control in this blot, such as GAPDH, beta-actin, ACTB, tubulin, or vinculin. "
            "Standard naming variants of the same control protein may be treated as the same value. "
            "Ignore control-role wording and non-identity details. "
            "Do not treat stains, dyes, total-protein stains, chemical reagents, or generic control labels as normalization-control proteins unless an actual control protein identity is explicitly given."
        ),

        "sample_type": (
            "The sample category used for blotting in this setup, such as cell lysate, whole cell lysate, tissue lysate, total protein extract, nuclear extract, cytoplasmic extract, serum, or plasma. "
            "Obvious near-equivalent wording within the same sample category may be treated as compatible, but distinct sample categories should not be merged. "
            "Ignore cell line names, tissue names, species names, disease labels, treatment groups, and preparation details when they do not change the sample-category meaning."
        ),

        "lysis_buffer": (
            "The lysis or extraction buffer identity used to prepare the sample, such as RIPA buffer, NP-40 lysis buffer, urea buffer, Laemmli buffer, or SDS lysis buffer. "
            "Standard naming variants of the same buffer may be treated as the same value. "
            "Ignore additives, inhibitors, salts, pH, vendor information, and procedural wording when they do not change the core buffer identity. "
            "Do not treat sample-type expressions such as cell lysate, tissue lysate, nuclear extract, serum, or plasma as lysis-buffer identity."
        ),

        "blocking_solution.identity": (
            "The blocking reagent identity or named blocking solution used before antibody incubation, such as BSA, bovine serum albumin, nonfat dry milk, skim milk, casein, normal serum, or a named blocking buffer product. "
            "Standard naming variants of the same blocking reagent may be treated as the same value. "
            "Ignore concentration, buffer context, incubation conditions, and vendor information unless they are part of the fixed name of a named blocking product. "
            "Do not treat detergents, wash buffers, antibody diluents, or other non-blocking reagents as blocking-solution identity."
        ),

        "blocking_solution.concentration": (
            "The concentration, dilution, or working-strength expression of the blocking solution, such as 5%, 3%, 1%, 1:10, or 1×. "
            "Use only values clearly tied to the blocking solution in the same setup. Equivalent surface variants that preserve the same quantitative meaning may be treated as compatible. "
            "Do not infer a blocking concentration when the paragraph only names the blocking reagent without giving an explicit concentration, dilution, or working strength."
        ),

        "gel_percentage": (
            "The explicit SDS-PAGE gel percentage or percentage range used in this blot, such as 8%, 10%, 12%, or 10-12%. "
            "Obvious surface variants may be treated as the same value. "
            "Use only values clearly presented as gel percentage. "
            "Do not use acrylamide composition details, buffer concentrations, or other nearby percentages unless they are explicitly the gel percentage."
        ),

        "loading_amount": (
            "The explicit sample or protein loading amount per lane or well, such as 20 μg, 30 μg, 50 μg/lane, or 10 μL/well. "
            "Equivalent unit-form variants that preserve the same loading amount may be treated as the same value. "
            "Use only values clearly presented as the loading amount for the blot. "
            "Do not use sample volume, antibody dilution, buffer volume, or other nearby quantities unless they are explicitly the loading amount."
        ),
    },

    "inorganic_materials": {
        "target_material": (
            "The target inorganic material system prepared or studied in this setup. Use the core material identity needed for stable distinction. "
            "Standard naming variants and formula versus full-name variants may be treated as the same value when they preserve the same material identity, but different compounds, phases, composites, solid solutions, and explicitly stated doped or substituted systems should not be merged. "
            "Do not rewrite a specific material system into a broader base material. "
            "Do not treat precursor names, additive names, morphology-only descriptors, or processing wording as target-material identity unless the paragraph clearly presents them as part of the target material itself."
        ),

        "precursors": (
            "The precursor or starting-material identity explicitly used to prepare the target material in this setup. Use the precursor identity itself. "
            "Standard naming variants and formula versus full-name variants may be treated as the same value when they preserve the same precursor identity, but chemically distinct precursors should not be merged. "
            "A setup may contain multiple explicit precursors; for a constraint on one precursor, mark match if that precursor is explicitly present, even if others are also present. "
            "Do not treat target-material wording, additive wording, or generic role labels as precursor identity unless a specific precursor is clearly supported."
        ),

        "dopant_or_additive.identity": (
            "The identity of a dopant, additive, sintering aid, or other composition-changing modifier explicitly introduced into this material system. Use the additive identity itself. "
            "Standard naming variants and formula versus full-name variants may be treated as the same value when they preserve the same additive identity, but chemically distinct additives should not be merged. "
            "A setup may contain multiple explicit additives; for a constraint on one additive, mark match if that additive is explicitly present, even if others are also present. "
            "Do not treat vague role labels, solvents, liquid media, or routine processing aids as additive identity unless the paragraph clearly supports them as meaningful additives in the setup."
        ),

        "dopant_or_additive.amount": (
            "The amount explicitly stated for a dopant or additive in this setup, such as wt%, mol%, at%, vol%, or ppm. "
            "Use only values clearly tied to the relevant dopant or additive in the same setup. Equivalent unit surface forms may be treated as the same value when they preserve the same meaning. "
            "Do not infer missing units or use values that are not clearly associated with the relevant additive."
        ),

        "synthesis_route": (
            "The main synthesis, consolidation, forming, or material-processing route used in this setup. "
            "This field may be matched either when the paragraph explicitly names the route, or when it clearly describes a standard and directly recognizable route through the sequence of experimental steps. "
            "A route does not need to be stated as a formal label if a knowledgeable reader can recover it from the setup description without guesswork. "
            "Use the route identity needed for stable distinction. Standard naming variants may be treated as the same value when they preserve the same route identity, but different synthesis families or clearly distinct route variants should not be merged. "
            "Do not infer a route from weak hints or scattered process details. Also do not treat a generic downstream step such as drying, debinding, or sintering alone as sufficient, unless the paragraph makes clear that this generic process itself is the main route."
        ),

        "atmosphere": (
            "The processing atmosphere explicitly used in this setup, such as air, vacuum, N2, Ar, H2, inert gas, or reducing atmosphere. "
            "Use the atmosphere identity itself rather than flow rate, pressure, temperature, or procedural wording. Standard naming variants may be treated as the same value when they preserve the same atmosphere meaning. "
            "Preserve distinctions that matter for matching, including different gases, mixed-gas atmospheres, vacuum versus non-vacuum conditions, and explicitly stated inert or reducing conditions. "
            "Do not infer atmosphere from pressure alone or from weak contextual hints."
        ),

        "heating_rate": (
            "The heating rate explicitly used in thermal processing in this setup, such as °C/min. "
            "Use only explicit heating-rate values. Equivalent unit surface forms may be treated as the same value when they preserve the same meaning. "
            "Do not use ramp descriptions without a specific rate value, and do not infer missing units or converted values that are not clearly supported."
        ),

        "pressing_pressure": (
            "The pressing or consolidation pressure explicitly used in this setup, such as MPa values tied to shaping, compaction, hot pressing, or another relevant consolidation step. "
            "Use only pressure values clearly tied to the relevant pressing or consolidation context. Equivalent unit surface forms may be treated as the same value when they preserve the same meaning. "
            "Do not treat gas pressure, chamber pressure, vacuum level, or force values that are not explicitly expressed as pressing pressure as matches."
        ),

        "sintering.temperature": (
            "The temperature used for the main sintering step in this setup. Use only explicit sintering-temperature values. "
            "Equivalent surface forms may be treated as the same value when they preserve the same temperature meaning. "
            "Do not treat temperatures from drying, calcination, annealing, debinding, melting, or other non-sintering steps as sintering temperature unless the paragraph clearly supports that they belong to the main sintering step."
        ),

        "sintering.time": (
            "The dwell or holding time used for the main sintering step in this setup. Use only explicit sintering-time values. "
            "Equivalent unit surface forms may be treated as the same value when they preserve the same time meaning. "
            "Do not treat times from drying, calcination, annealing, debinding, or other non-sintering steps as sintering time unless the paragraph clearly supports that they belong to the main sintering step."
        ),

        "post_treatment.type": (
            "The type of distinct post-treatment step performed before or after the main sintering step in this setup, such as annealing, calcination, oxidation, reduction, quenching, or heat treatment. "
            "Use the treatment identity needed for stable distinction. Standard naming variants may be treated as the same value when they preserve the same treatment identity, but different post-treatment families or clearly distinct variants should not be merged. "
            "Do not treat routine drying, cooling, storage, or mechanical surface treatment as post-treatment unless the paragraph clearly presents them as meaningful post-treatment steps."
        ),

        "post_treatment.temperature": (
            "The temperature used for a distinct post-treatment step in this setup. Use only explicit post-treatment temperature values. "
            "Equivalent surface forms may be treated as the same value when they preserve the same temperature meaning. "
            "Do not treat temperatures from the main sintering step, drying, debinding, or other non-post-treatment steps as post-treatment temperature unless the paragraph clearly supports the same post-treatment context."
        ),

        "post_treatment.time": (
            "The time used for a distinct post-treatment step in this setup. Use only explicit post-treatment time values. "
            "Equivalent unit surface forms may be treated as the same value when they preserve the same time meaning. "
            "Do not treat times from the main sintering step or from other non-post-treatment steps as post-treatment time unless the paragraph clearly supports the same post-treatment context."
        )
    },
   
    "organic_synthesis": {
        "transformation_type": (
            "The reaction class or transformation type carried out in this setup, such as Suzuki coupling, Sonogashira coupling, hydrogenation, oxidation, reduction, cyclization, cycloaddition, amidation, arylation, or borylation. "
            "Use the most specific transformation identity clearly supported by the paragraph. Standard naming variants may be treated as the same value when they preserve the same transformation identity, but distinct transformation types should not be merged. "
            "Short, evidence-grounded chemical reasoning is allowed only when the paragraph provides a clear and diagnostic reaction pattern. "
            "Do not assign a transformation type from weak, generic, or ambiguous reaction cues alone."
        ),

        "substrates": (
            "The key stoichiometric starting materials or reaction partners transformed in this setup. Use the substrate or reaction-partner identity itself, not amounts, equivalents, local paper-specific labels, or procedural wording. "
            "Standard naming variants may be treated as the same value when they preserve the same chemical identity, but chemically distinct substrates should not be merged. "
            "A setup may contain multiple explicit substrates; for a constraint on one substrate, mark match if that substrate is explicitly present as a transformed reaction partner. "
            "Do not treat catalysts, ligands, bases, additives, solvents, products, intermediates, or routine reagents as substrates unless the paragraph clearly presents them as reaction partners being transformed."
        ),

        "catalysts.name": (
            "The catalyst identity used in the reaction, such as Pd(PPh3)4, PdCl2(PPh3)2, CuTC, or RuCl3. "
            "Use the catalyst identity itself and ignore loading, concentration, supplier information, and other non-identity details. Standard naming variants of the same catalyst may be treated as the same value. "
            "Do not treat free ligands, bases, additives, solvents, substrates, or other non-catalyst roles as catalyst identity."
        ),

        "catalysts.loading": (
            "The catalyst loading explicitly stated for the reaction, such as mol%, equiv, mol, mmol, or μmol when clearly tied to the catalyst. "
            "Use only values clearly associated with the catalyst in the same setup. Equivalent surface forms may be treated as the same value when they preserve the same quantitative meaning. "
            "Do not infer, recalculate, or convert values beyond what is directly supported. "
            "Do not use amounts belonging to substrates, ligands, bases, additives, or solvents as catalyst loading."
        ),

        "ligands": (
            "The ligand identity used in the reaction, such as PPh3, BINAP, bpy, or XPhos. "
            "Standard naming variants of the same ligand may be treated as the same value. A ligand may also be matched when it is directly recoverable from a named catalyst complex and the paragraph clearly supports that relation. "
            "Do not treat metal names, whole catalyst identities, bases, additives, solvents, substrates, or local paper-specific codes alone as ligand identity."
        ),

        "bases": (
            "The base identity used in the reaction, such as Cs2CO3, Na2CO3, K3PO4, DBU, DIPEA, or Et3N. "
            "Standard naming variants of the same base may be treated as the same value. Ignore amount, concentration, and other non-identity details. "
            "Do not treat catalysts, ligands, additives, solvents, substrates, or general reagents as bases unless the paragraph clearly supports a base role."
        ),

        "additives": (
            "The identity of a functional additive or promoter used in the reaction, such as an oxidant, reductant, acid, salt, or other explicitly stated additive with a distinct reaction role. "
            "Standard naming variants of the same additive may be treated as the same value when they preserve the same chemical identity and role. Ignore amount, concentration, and other non-identity details. "
            "Do not treat catalysts, ligands, bases, solvents, substrates, or workup reagents as additives. "
            "A reagent may be treated as an additive even when used stoichiometrically if the paragraph clearly assigns it a functional additive role."
        ),

        "solvents": (
            "The solvent or solvent components used as the reaction medium, such as DME, toluene, THF, water, or H2O. "
            "Use the solvent identity itself and ignore volume, ratio, drying status, degassing status, and other non-identity details. Standard naming variants of the same solvent may be treated as the same value. "
            "A setup may contain multiple explicit solvent components; for a constraint on one solvent, mark match if that solvent is explicitly present in the reaction medium. "
            "Do not infer solvent identity from vague solution descriptions alone, and do not treat catalysts, ligands, bases, additives, substrates, electrolytes, anolytes, or catholytes as solvents."
        ),

        "activation_mode": (
            "The high-level reaction-driving or activation mode used in the setup, such as thermal, photochemical, electrochemical, microwave, mechanochemical, sonochemical, or plasma. "
            "Use a canonical activation label only when the paragraph clearly supports it. Obvious local evidence may support these labels, such as heating or reflux for thermal, light irradiation or LED for photochemical, electrolysis or applied current for electrochemical, microwave irradiation for microwave, ball milling or grinding for mechanochemical, ultrasound or sonication for sonochemical, and plasma discharge for plasma. "
            "Do not infer activation mode from weak or ambiguous procedural context alone."
        )
    },
    
    "polymer_materials": {
        "polymer_system": (
            "The main photocurable polymer, resin, monomer, oligomer, or crosslinkable network system used in this curing setup. Use the core polymer-system identity needed for stable distinction. "
            "Ignore concentration, ratio, molecular-weight grade, supplier information, and surrounding formulation wording when they do not change the core identity. "
            "Obvious naming variants of the same polymer system may be treated as the same value, but distinct polymer systems should not be merged. "
            "Do not treat photoinitiators, crosslinking agents, solvents, fillers, nanoparticles, proteins, peptides, drugs, or other additives as the main polymer system unless the paragraph clearly presents them as part of the polymer-system identity itself."
        ),

        "polymer_concentration": (
            "The concentration, fraction, loading, or amount-per-volume value explicitly tied to the main polymer system in this curing setup. "
            "Use only values clearly associated with the main polymer system rather than with photoinitiators, crosslinking agents, solvents, or other additives. "
            "Equivalent concentration forms may be treated as compatible when they preserve the same underlying polymer concentration meaning. "
            "Do not infer missing units or accept values that require assumptions beyond the explicit paragraph evidence."
        ),

        "photoinitiators.name": (
            "The identity of the photoinitiator used in the curing setup, such as LAP, Irgacure 2959, TPO, BAPO, eosin Y, or camphorquinone. "
            "Use the initiator identity itself and ignore loading, supplier information, solvents, and surrounding formulation wording when they do not change the identity. "
            "Obvious naming variants of the same photoinitiator may be treated as the same value, but distinct initiators should not be merged. "
            "Do not treat polymers, crosslinking agents, co-monomers, solvents, fillers, nanoparticles, or generic phrases such as photoinitiator system as photoinitiator identity unless a specific initiator is clearly supported."
        ),

        "photoinitiators.loading": (
            "The loading or concentration explicitly stated for the photoinitiator in this curing setup, such as wt%, mol%, % w/v, or mg/mL. "
            "Use only values clearly tied to a photoinitiator rather than to the main polymer system, a crosslinking agent, or another additive. "
            "Equivalent numeric forms may be treated as compatible when they preserve the same loading meaning. "
            "Do not infer missing units or accept values that require extra assumptions."
        ),

        "crosslinking_agent.name": (
            "The identity of an explicitly added crosslinking agent or other distinct reactive formulation additive used together with the main polymer system. "
            "Use the component identity itself and ignore loading, supplier information, and surrounding formulation wording when they do not change the identity. "
            "Obvious naming variants of the same component may be treated as the same value, but distinct reactive additives should not be merged. "
            "Do not treat the main polymer system, photoinitiators, solvents, fillers, nanoparticles, catalysts, inhibitors, or generic phrases such as crosslinker system as crosslinking-agent identity unless the paragraph clearly supports that role."
        ),

        "crosslinking_agent.loading": (
            "The loading or concentration explicitly stated for a crosslinking agent or other retained reactive crosslinking-related additive in this curing setup. "
            "Use only values clearly tied to the relevant crosslinking agent rather than to the main polymer system, a photoinitiator, or another additive. "
            "Equivalent numeric forms may be treated as compatible when they preserve the same loading meaning. "
            "Do not infer missing units or accept values that require extra assumptions."
        ),

        "wavelength": (
            "The irradiation wavelength explicitly used in the curing setup, such as 365 nm, 405 nm, or a narrow wavelength range. "
            "Use only explicit wavelength values or narrow wavelength ranges. "
            "Ignore light-source names, color descriptions, intensity, exposure time, and other non-wavelength details. "
            "Do not infer wavelength from generic descriptions such as UV, visible, blue light, or near-IR alone."
        ),

        "intensity": (
            "The light intensity or irradiance explicitly used in the curing setup. Use only values clearly tied to the curing condition. "
            "Ignore wavelength, exposure time, light-source type, and surrounding procedural wording when they do not change the intensity meaning. "
            "Equivalent numeric forms may be treated as compatible when they preserve the same intensity meaning. "
            "Do not infer missing units or accept values that require extra assumptions."
        ),

        "exposure_time": (
            "The irradiation or exposure time explicitly used in the curing setup. Use only values clearly tied to the curing condition. "
            "Ignore wavelength, intensity, light-source type, and surrounding procedural wording when they do not change the exposure-time meaning. "
            "Equivalent numeric forms may be treated as compatible when they preserve the same time meaning. "
            "Do not infer missing units or accept values that require extra assumptions."
        ),

        "light_source_type": (
            "The type of light source used for curing, such as LED, UV LED, laser, projector, mercury lamp, xenon lamp, halogen lamp, or UV lamp. "
            "Use the source identity itself rather than wavelength, intensity, exposure time, model number, or supplier information. Prefer the most specific supported source identity. "
            "Obvious naming variants of the same source type may be treated as the same value, but distinct source types should not be merged. "
            "Do not treat generic phrases such as light, irradiation, or photo-curing as light-source type when no specific source is given."
        )
    },
}





def to_prompt_json(obj: Any) -> str:
    return json.dumps(obj, ensure_ascii=False, indent=2)


def collect_relevant_fields(constraints: List[Dict[str, Any]], explore_field: Optional[str] = None) -> List[str]:
    fields: List[str] = []
    for c in constraints or []:
        field = c.get("field")
        if field and field not in fields:
            fields.append(field)
    if explore_field and explore_field not in fields:
        fields.append(explore_field)
    return fields


def get_relevant_field_definitions(
    domain: str,
    constraints: List[Dict[str, Any]],
    explore_field: Optional[str] = None,
) -> Dict[str, str]:
    if domain not in JUDGE_FIELD_DEFINITIONS:
        raise KeyError(f"Unknown domain: {domain}")
    field_definitions = JUDGE_FIELD_DEFINITIONS[domain]
    relevant_fields = collect_relevant_fields(constraints, explore_field)
    return {f: field_definitions[f] for f in relevant_fields if f in field_definitions}



def build_reproduction_prompt(
    domain: str,
    constraints: List[Dict[str, Any]],
    paragraph_text: str,
) -> str:
    relevant_defs = get_relevant_field_definitions(domain, constraints, None)
    instance_block = REPRODUCTION_PROMPT.format(
        constraints_json=to_prompt_json(constraints),
        field_definitions_json=to_prompt_json(relevant_defs),
        paragraph_text=paragraph_text,
    )
    return instance_block


def build_exploration_prompt(
    domain: str,
    constraints: List[Dict[str, Any]],
    explore_field: str,
    paragraph_text: str,
) -> str:
    relevant_defs = get_relevant_field_definitions(domain, constraints, explore_field)
    instance_block = EXPLORATION_PROMPT.format(
        explore_field=explore_field,
        constraints_json=to_prompt_json(constraints),
        field_definitions_json=to_prompt_json(relevant_defs),
        paragraph_text=paragraph_text,
    )
    return instance_block



if __name__ == "__main__":
    
    OUTPUT_DIR = Path("data/manual_audit")


    def style_worksheet(ws):
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(name="Times New Roman", bold=True, color="FFFFFF")
        body_font = Font(name="Times New Roman")
        thin = Side(style="thin", color="D9D9D9")

        # 表头样式
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)

        # 内容样式
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:D{ws.max_row}"

        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 38
        ws.column_dimensions["C"].width = 110



    def export_domain_field_definitions(judge_field_definitions: dict, output_dir: Path):
        output_dir.mkdir(parents=True, exist_ok=True)

        for domain, field_map in judge_field_definitions.items():
            wb = Workbook()
            ws = wb.active
            ws.title = "field_definitions"

            # 表头
            ws.append(["domain", "field_name", "definition"])

            # 数据
            for field_name, definition in sorted(field_map.items()):
                ws.append([domain, field_name, definition, ""])

            style_worksheet(ws)

            out_path = output_dir / f"field_def_{domain}.xlsx"
            wb.save(out_path)
            print(f"Saved: {out_path}")


    export_domain_field_definitions(JUDGE_FIELD_DEFINITIONS, OUTPUT_DIR)

# python scripts/relevance/prompt.py